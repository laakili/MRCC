import pandas as pd
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from app.models.personnel import Personnel
from app.models.absence import Absence
from app.models.formation import Formation, InscriptionFormation
from app.models.mission import Mission
from datetime import datetime

# ==================== EXPORT PERSONNEL ====================

def export_personnel_excel(grade_id=None, unite_id=None):
    """Exporte la liste du personnel au format Excel"""
    
    # Récupérer les données
    query = Personnel.query
    
    if grade_id:
        query = query.filter_by(grade_id=grade_id)
    if unite_id:
        query = query.filter_by(unite_orga_id=unite_id)
    
    personnel = query.all()
    
    # Créer un DataFrame pandas
    data = []
    for p in personnel:
        data.append({
            'Matricule': p.matricule,
            'Nom': p.nom,
            'Prénom': p.prenom,
            'CIN': p.carte_identite,
            'Date Naissance': p.date_de_naissance.strftime('%d/%m/%Y') if p.date_de_naissance else '',
            'Grade': p.grade_ref.libelle if p.grade_ref else '',
            'Emploi': p.emploi_ref.identifiant_emploi if p.emploi_ref else '',
            'Unité': p.unite_ref.unite_organisat if p.unite_ref else '',
            'Date Recrutement': p.date_recrutement.strftime('%d/%m/%Y') if p.date_recrutement else '',
            'Statut': p.statut,
            'Téléphone': p.telephone,
            'Email': p.email_professionnel
        })
    
    df = pd.DataFrame(data)
    
    # Créer un fichier Excel en mémoire
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Personnel', index=False)
        
        # Formater le fichier Excel
        workbook = writer.book
        worksheet = writer.sheets['Personnel']
        
        # En-têtes en gras
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
            cell.font = Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center')
        
        # Ajuster la largeur des colonnes
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
    
    output.seek(0)
    return output


# ==================== EXPORT ABSENCES ====================

def export_absences_excel(type_absence=None, statut=None, date_debut=None, date_fin=None):
    """Exporte les absences au format Excel"""
    
    query = Absence.query
    
    if type_absence:
        query = query.filter_by(type_absence=type_absence)
    if statut:
        query = query.filter_by(statut=statut)
    if date_debut:
        query = query.filter(Absence.date_debut >= datetime.strptime(date_debut, '%Y-%m-%d'))
    if date_fin:
        query = query.filter(Absence.date_fin <= datetime.strptime(date_fin, '%Y-%m-%d'))
    
    absences = query.all()
    
    data = []
    for a in absences:
        data.append({
            'Matricule': a.personnel.matricule,
            'Nom': a.personnel.nom,
            'Prénom': a.personnel.prenom,
            'Type': a.type_absence,
            'Date début': a.date_debut.strftime('%d/%m/%Y'),
            'Date fin': a.date_fin.strftime('%d/%m/%Y'),
            'Durée (jours)': float(a.duree_jours),
            'Statut': a.statut,
            'Motif': a.motif
        })
    
    df = pd.DataFrame(data)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Absences', index=False)
    
    output.seek(0)
    return output


# ==================== EXPORT FORMATIONS ====================

def export_formations_excel(annee=None, statut=None, type_formation=None):
    """Exporte la liste des formations au format Excel"""
    
    query = Formation.query
    
    if annee:
        query = query.filter(db.extract('year', Formation.date_debut_prevue) == annee)
    if statut:
        query = query.filter_by(statut=statut)
    if type_formation:
        query = query.filter_by(type_formation=type_formation)
    
    formations = query.all()
    
    data = []
    for f in formations:
        data.append({
            'Code': f.code_formation,
            'Intitulé': f.intitule,
            'Type': f.type_formation,
            'Catégorie': f.categorie,
            'Organisme': f.organisme,
            'Date début': f.date_debut_prevue.strftime('%d/%m/%Y') if f.date_debut_prevue else '',
            'Date fin': f.date_fin_prevue.strftime('%d/%m/%Y') if f.date_fin_prevue else '',
            'Durée (jours)': f.duree_jours,
            'Coût estimé': float(f.cout_estime) if f.cout_estime else 0,
            'Statut': f.statut,
            'Inscrits': f.nb_inscrits
        })
    
    df = pd.DataFrame(data)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Formations', index=False)
    
    output.seek(0)
    return output



def export_participants_formation_excel(formation_id):
    """Exporte la liste des participants à une formation"""
    
    inscriptions = InscriptionFormation.query.filter_by(formation_id=formation_id).all()
    
    data = []
    for ins in inscriptions:
        data.append({
            'Matricule': ins.personnel.matricule,
            'Nom': ins.personnel.nom,
            'Prénom': ins.personnel.prenom,
            'Unité': ins.personnel.unite_ref.unite_organisat if ins.personnel.unite_ref else '',
            'Statut inscription': ins.statut,
            'Présent': 'Oui' if ins.present else 'Non',
            'Note': float(ins.note_obtenue) if ins.note_obtenue else '',
            'Certificat': 'Oui' if ins.certificat_obtenu else 'Non'
        })
    
    df = pd.DataFrame(data)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Participants', index=False)
    
    output.seek(0)
    return output


# ==================== EXPORT MISSIONS ====================

def export_missions_excel(annee=None, statut=None, nature_id=None):
    """Exporte la liste des missions au format Excel"""
    
    query = Mission.query
    
    if annee:
        query = query.filter_by(annee=annee)
    if statut:
        query = query.filter_by(statut=statut)
    if nature_id:
        query = query.filter_by(nature_id=nature_id)
    
    missions = query.all()
    
    data = []
    for m in missions:
        data.append({
            'N° Ordre': m.numero_ordre_mission,
            'Agent': f"{m.personnel.nom} {m.personnel.prenom}",
            'Thème': m.theme,
            'Nature': m.nature.libelle if m.nature else '',
            'Lieu': m.lieu_mission,
            'Date début': m.date_debut.strftime('%d/%m/%Y'),
            'Date fin': m.date_fin.strftime('%d/%m/%Y') if m.date_fin else '',
            'Durée': m.duree_prevue,
            'Montant': float(m.montant_estime) if m.montant_estime else 0,
            'Statut': m.statut
        })
    
    df = pd.DataFrame(data)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Missions', index=False)
    
    output.seek(0)
    return output

# ==================== EXPORT FORMATIONS PDF ====================

def export_plan_formation_pdf(formation):
    """Exporte le plan de formation au format PDF avec ReportLab"""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import io
    from datetime import datetime
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []
    
    # Titre
    elements.append(Paragraph("OFFICE NATIONAL DES AÉROPORTS", styles['Title']))
    elements.append(Spacer(1, 0.5*cm))
    elements.append(Paragraph(f"PLAN DE FORMATION", styles['Heading1']))
    elements.append(Spacer(1, 0.5*cm))
    
    # Informations de la formation
    elements.append(Paragraph(f"{formation.code_formation} - {formation.intitule}", styles['Heading2']))
    elements.append(Spacer(1, 0.5*cm))
    
    # Détails de la formation
    data = [
        ['Type:', formation.type_formation],
        ['Catégorie:', formation.categorie],
        ['Niveau:', formation.niveau or '-'],
        ['Dates:', f"{formation.date_debut_prevue.strftime('%d/%m/%Y') if formation.date_debut_prevue else '-'} au {formation.date_fin_prevue.strftime('%d/%m/%Y') if formation.date_fin_prevue else '-'}"],
        ['Durée:', f"{formation.duree_jours} jours ({formation.duree_heures} heures)"],
        ['Lieu:', formation.lieu or '-'],
        ['Organisme:', formation.organisme or '-'],
        ['Formateur:', formation.formateur_nom or '-'],
    ]
    
    table = Table(data, colWidths=[4*cm, 12*cm])
    table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('GRID', (0, 0), (-1, -1), 1, colors.lightgrey),
        ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.5*cm))
    
    # Objectifs
    elements.append(Paragraph("Objectifs", styles['Heading3']))
    elements.append(Paragraph(formation.objectifs or 'Non spécifiés', styles['Normal']))
    elements.append(Spacer(1, 0.5*cm))
    
    # Programme
    elements.append(Paragraph("Programme", styles['Heading3']))
    elements.append(Paragraph(formation.programme or 'Non spécifié', styles['Normal']))
    elements.append(Spacer(1, 0.5*cm))
    
    # Liste des participants
    elements.append(Paragraph(f"Participants inscrits ({formation.nb_inscrits})", styles['Heading3']))
    elements.append(Spacer(1, 0.3*cm))
    
    # Tableau des participants
    participants_data = [['Matricule', 'Nom', 'Prénom', 'Statut']]
    for inscrit in formation.inscriptions:
        participants_data.append([
            inscrit.personnel.matricule,
            inscrit.personnel.nom,
            inscrit.personnel.prenom,
            inscrit.statut
        ])
    
    if len(participants_data) > 1:
        table_participants = Table(participants_data, colWidths=[3*cm, 4*cm, 4*cm, 3*cm])
        table_participants.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ]))
        elements.append(table_participants)
    else:
        elements.append(Paragraph("Aucun participant inscrit", styles['Normal']))
    
    # Pied de page
    elements.append(Spacer(1, 2*cm))
    elements.append(Paragraph(f"Document généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}", styles['Normal']))
    
    doc.build(elements)
    pdf = buffer.getvalue()
    buffer.close()
    return pdf


def export_plan_formation_pdf1(formation):
    """
    Exporte le plan de formation au format PDF
    """
    from weasyprint import HTML
    from flask import render_template
    
    try:
        # Essayer d'utiliser le template HTML
        html = render_template('pdf/plan_formation.html', formation=formation, now=datetime.now)
        pdf = HTML(string=html).write_pdf()
        return pdf
    except Exception as e:
        print(f"Erreur lors de la génération du PDF avec template: {e}")
        
        # Fallback: générer un PDF simple avec les informations de base
        try:
            html = generate_simple_formation_html(formation)
            pdf = HTML(string=html).write_pdf()
            return pdf
        except Exception as e2:
            print(f"Erreur lors de la génération du PDF simple: {e2}")
            return generate_error_pdf(f"Erreur: {str(e)}")


def export_participants_formation_excel(formation_id):
    """Exporte la liste des participants à une formation au format Excel"""
    from app.models.formation import InscriptionFormation, Formation
    from app.models.personnel import Personnel, UniteOrganisationnelle
    import pandas as pd
    import io
    
    # Récupérer les inscriptions avec les relations
    inscriptions = InscriptionFormation.query.filter_by(formation_id=formation_id)\
        .join(Personnel)\
        .options(db.joinedload(InscriptionFormation.personnel).joinedload(Personnel.unite_ref))\
        .all()
    
    formation = Formation.query.get(formation_id)
    
    data = []
    for ins in inscriptions:
        data.append({
            'Matricule': ins.personnel.matricule,
            'Nom': ins.personnel.nom,
            'Prénom': ins.personnel.prenom,
            'Unité': ins.personnel.unite_ref.unite_organisat if ins.personnel.unite_ref else '',
            'Statut inscription': ins.statut,
            'Présent': 'Oui' if ins.present else 'Non',
            'Note': float(ins.note_obtenue) if ins.note_obtenue else '',
            'Certificat': 'Oui' if ins.certificat_obtenu else 'Non',
            'Date inscription': ins.date_inscription.strftime('%d/%m/%Y')
        })
    
    df = pd.DataFrame(data)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name=f'Participants_{formation.code_formation}'[:31], index=False)
        
        # Formatage
        workbook = writer.book
        worksheet = writer.sheets[f'Participants_{formation.code_formation}'[:31]]
        
        # Ajuster la largeur des colonnes
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
    
    output.seek(0)
    return output

def generate_error_pdf(error_message):
    """Génère un PDF simple en cas d'erreur"""
    from weasyprint import HTML
    
    html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="UTF-8">
        <title>Erreur</title>
        <style>
            body {{ 
                font-family: Arial, sans-serif; 
                margin: 50px; 
                line-height: 1.6;
            }}
            .error {{ 
                color: #721c24; 
                background-color: #f8d7da;
                border: 1px solid #f5c6cb;
                padding: 20px; 
                border-radius: 5px;
                margin-top: 20px;
            }}
            h1 {{ color: #dc3545; }}
        </style>
    </head>
    <body>
        <h1>Erreur de génération PDF</h1>
        <div class="error">{error_message}</div>
        <p>Veuillez contacter l'administrateur système.</p>
        <hr>
        <p><small>ONDA - Système de Gestion RH</small></p>
    </body>
    </html>
    """
    return HTML(string=html).write_pdf()


def generate_participants_excel(inscriptions):
    """Génère un fichier Excel avec la liste de tous les participants"""
    import pandas as pd
    import io
    from openpyxl.styles import Font, Alignment, PatternFill
    
    data = []
    for ins in inscriptions:
        data.append({
            'Date inscription': ins.date_inscription.strftime('%d/%m/%Y %H:%M'),
            'Matricule': ins.personnel.matricule,
            'Nom': ins.personnel.nom,
            'Prénom': ins.personnel.prenom,
            'Unité': ins.personnel.unite_ref.unite_organisat if ins.personnel.unite_ref else '',
            'Grade': ins.personnel.grade_ref.libelle if ins.personnel.grade_ref else '',
            'Code Formation': ins.formation.code_formation,
            'Formation': ins.formation.intitule,
            'Type formation': ins.formation.type_formation,
            'Catégorie': ins.formation.categorie,
            'Date début formation': ins.formation.date_debut_prevue.strftime('%d/%m/%Y') if ins.formation.date_debut_prevue else '',
            'Date fin formation': ins.formation.date_fin_prevue.strftime('%d/%m/%Y') if ins.formation.date_fin_prevue else '',
            'Statut inscription': ins.statut,
            'Présent': 'Oui' if ins.present else 'Non',
            'Note': float(ins.note_obtenue) if ins.note_obtenue else '',
            'Certificat obtenu': 'Oui' if ins.certificat_obtenu else 'Non',
            'Priorité': ins.priorite,
            'Motivation': ins.motivation or '',
            'Validé par': ins.validateur.email if ins.validateur else '',
            'Date validation': ins.date_validation.strftime('%d/%m/%Y %H:%M') if ins.date_validation else '',
        })
    
    df = pd.DataFrame(data)
    
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Tous les participants', index=False)
        
        # Formatage
        workbook = writer.book
        worksheet = writer.sheets['Tous les participants']
        
        # En-têtes en gras et colorés
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='003366', end_color='003366', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        # Ajuster la largeur des colonnes
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
    
    output.seek(0)
    return output
